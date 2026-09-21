"""CargoVeritas shipping-document verification pipeline."""
from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path

from loader import GROSS_WEIGHT_PATTERN, Inbox, extract_two_column_fields

FIELDS = ("shipper", "consignee", "notify_party", "port_of_loading", "port_of_discharge", "container_count", "gross_weight_kg")


def classify_email(subject: str, body: str, attachment_names=()) -> str:
    """Route every email before any document work is attempted."""
    text = f"{subject}\n{body}".lower()
    names = " ".join(attachment_names).lower()
    has_si = bool(re.search(r"(?:^|[^a-z])si(?:[^a-z]|$)", names)) or "shipping instruction" in names
    has_bl = bool(re.search(r"(?:^|[^a-z])bl(?:[^a-z]|$)", names)) or "bill of lading" in names or "draft bl" in names
    if (has_si and has_bl) or ("draft" in text and (" bill of lading" in text or " bl" in text) and any(w in text for w in ("compare", "check", "confirm"))):
        return "BL_COMPARISON"
    if any(word in text for word in ("unsubscribe", "limited offer", "buy now", "promotion", "marketing")):
        return "SPAM"
    if any(word in text for word in ("invoice", "payment", "billing", "charge", "freight rate")):
        return "INVOICE_QUERY"
    if any(word in text for word in ("shipping instruction", "submit si", "amend si", "prepare si")):
        return "SI_REQUEST"
    return "GENERAL"


def extract_text_from_bytes(raw: bytes, filename: str) -> str:
    """Extract readable content from common logistics attachment formats."""
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in (".txt", ".csv"):
            try:
                return raw.decode("utf-8")
            except UnicodeDecodeError:
                return raw.decode("latin-1", errors="ignore")
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw), strict=False)
            text = "\n".join(page.extract_text(extraction_mode="layout") or page.extract_text() or "" for page in reader.pages).strip()
            if text:
                return text
            # Some carrier PDFs, including the supplied email_514 documents,
            # are a scanned page image with no selectable text layer. OCR only
            # runs for that case; normal text PDFs stay on the fast path above.
            try:
                import numpy as np
                from rapidocr_onnxruntime import RapidOCR
                engine = RapidOCR()
                lines = []
                for page in reader.pages:
                    for image in page.images:
                        result, _ = engine(np.asarray(image.image.convert("RGB")))
                        lines.extend(item[1] for item in (result or []) if len(item) > 1 and item[1])
                return "\n".join(lines).strip()
            except Exception:
                return ""
        if suffix in (".docx", ".doc"):
            import docx
            doc = docx.Document(io.BytesIO(raw))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))
            return "\n".join(paragraphs).strip()
        if suffix in (".xlsx", ".xls"):
            import openpyxl
            book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            rows = []
            for sheet in book.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if cells:
                        rows.append(" | ".join(cells))
            return "\n".join(rows).strip()
    except Exception:
        # DOCX and XLSX are ZIP-based XML. This fallback still recovers text
        # when an office library rejects a slightly malformed export.
        if suffix in (".docx", ".xlsx"):
            try:
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    names = archive.namelist()
                    xml_files = (["word/document.xml"] if suffix == ".docx" else [name for name in names if name.startswith("xl/sharedStrings") or name.startswith("xl/worksheets/")])
                    parts = []
                    for name in xml_files:
                        if name in names:
                            xml = archive.read(name).decode("utf-8", errors="ignore")
                            parts.append(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", xml)))
                    return "\n".join(part.strip() for part in parts if part.strip())
            except Exception:
                pass
        return ""
    return ""


def _label_value(text: str, labels) -> str | None:
    for label in labels:
        # PDFs preserve label/value columns as multiple spaces; Word and Excel
        # tables arrive as pipe-delimited text after extraction.
        match = re.search(r"(?im)^\s*" + label + r"(?:\s*(?::|\-|\|)\s*|\s{2,})(.+?)\s*$", text)
        if match:
            return match.group(1).strip(" |\t")
    return None


def extract_fields_from_doc(text: str) -> dict:
    """Extract exactly the seven mandatory values; unknown values remain null."""
    # Bilingual carrier templates add one or more parenthetical translations to
    # field labels, e.g. "POD (\u5378\u8d27\u6e2f)". Remove label annotations before matching.
    label_terms = r"shipper(?:/exporter)?|consignee|notify(?:\s+party)?|port\s*of\s*loading|load(?:ing)?\s*port|pol|port\s*of\s*discharge|discharge\s*port|pod|(?:no\.\s*of\s*)?containers?(?:\s*or\s*packages)?|container\s*count|total\s*containers|(?:total\s*)?gross\s*(?:weight|wt)"
    text = re.sub(r"(?im)(" + label_terms + r")(?:\s*[\(\uff08][^\)\uff09]*[\)\uff09])+", r"\1", text)
    result = {field: None for field in FIELDS}
    result.update(extract_two_column_fields(text))
    result["shipper"] = result["shipper"] or _label_value(text, (r"shipper(?:/exporter)?(?:\s*\([^)]*\))?",))
    result["consignee"] = result["consignee"] or _label_value(text, (r"consignee(?:\s*\([^)]*\))?",))
    result["notify_party"] = result["notify_party"] or _label_value(text, (r"notify(?:\s+party)?",))
    result["port_of_loading"] = result["port_of_loading"] or _label_value(text, (r"port\s*of\s*loading(?:\s*\(pol\))?", r"load(?:ing)?\s*port", r"pol"))
    result["port_of_discharge"] = result["port_of_discharge"] or _label_value(text, (r"port\s*of\s*discharge(?:\s*\(pod\))?", r"discharge\s*port", r"pod"))
    containers = result["container_count"] or _label_value(text, (r"(?:no\.\s+of\s+)?containers?(?:\s+or\s+packages)?", r"container\s+count", r"total\s+containers"))
    weight = result["gross_weight_kg"] or _label_value(text, (r"(?:total\s+)?gross\s*(?:weight|wt)(?:\s*\(kgs?\))?",))
    if not weight:
        fallback = GROSS_WEIGHT_PATTERN.search(text)
        weight = fallback.group(1) if fallback else None
    if containers:
        number = re.search(r"\d+", containers.replace(",", ""))
        result["container_count"] = int(number.group()) if number else None
    if weight:
        number = re.search(r"\d[\d,]*(?:\.\d+)?", weight)
        result["gross_weight_kg"] = float(number.group().replace(",", "")) if number else None
    return result


def _normal_text(value) -> str:
    return " ".join(re.sub(r"[^a-z0-9]", " ", str(value or "").lower()).split())


def _normal_port(value) -> str:
    text = _normal_text(value)
    aliases = {"port klang": "mypkg", "pkg": "mypkg", "mypkg": "mypkg", "callao": "pecll", "pecll": "pecll", "singapore": "sgsin", "sg sin": "sgsin", "sgsin": "sgsin", "karachi": "pkkhi", "pk khi": "pkkhi", "pkkhi": "pkkhi"}
    return next((code for name, code in aliases.items() if name in text), text)


def _primary_entity(value) -> str:
    """Return the primary company name, without trailing address details."""
    first = re.split(r"[|;\r\n]", str(value or ""), maxsplit=1)[0]
    return _normal_text(first)


def _same_entity(left, right) -> bool:
    left, right = _primary_entity(left), _primary_entity(right)
    if not left or not right:
        return False
    # Never hide a primary-brand typo behind a broad fuzzy/company match.
    if left.split()[0] != right.split()[0]:
        return False
    def subsidiary(value):
        return bool(re.search(r"\b(?:middle\s+east|fze|fzc|regional|subsidiary)\b", value))
    if subsidiary(left) != subsidiary(right):
        return False
    # Entity fields often differ only because one document retains an address.
    return left == right or (min(len(left), len(right)) >= 5 and (left in right or right in left))


def _normal_number(value, *, integer=False):
    match = re.search(r"\d[\d,\s]*(?:\.\d+)?", str(value or ""))
    if not match:
        return None
    number = float(re.sub(r"[,\s]", "", match.group()))
    return int(number) if integer else number


def compare_fields(si_data: dict, bl_data: dict) -> list[str]:
    """Compare all fields after normalizing harmless formatting differences."""
    defects = []
    for field in FIELDS:
        si, bl = si_data.get(field), bl_data.get(field)
        if field == "container_count":
            same = _normal_number(si, integer=True) == _normal_number(bl, integer=True)
        elif field == "gross_weight_kg":
            si_number, bl_number = _normal_number(si), _normal_number(bl)
            same = si_number is not None and bl_number is not None and abs(si_number - bl_number) < 0.01
        elif field in ("port_of_loading", "port_of_discharge"):
            same = _normal_port(si) == _normal_port(bl)
        elif field in ("shipper", "consignee", "notify_party"):
            same = _same_entity(si, bl)
        else:
            same = _normal_text(si) == _normal_text(bl)
        if not same:
            defects.append(field)
    return defects


def process_email(email: dict, attachment_texts: dict[str, str]) -> dict:
    """Return the structured category, extraction, comparison, and review record."""
    attachments = list(attachment_texts)
    category = classify_email(email.get("subject", ""), email.get("body", ""), attachments)
    result = {"category": category, "status": "AUTO_RESOLVED", "review_reason": None, "defect_fields": [], "si_data": {}, "bl_data": {}}
    if category != "BL_COMPARISON":
        return result
    si_name = next((name for name in attachments if re.search(r"(?:^|[^a-z])si(?:[^a-z]|$)", name.lower()) or "shipping instruction" in name.lower()), None)
    bl_name = next((name for name in attachments if re.search(r"(?:^|[^a-z])bl(?:[^a-z]|$)", name.lower()) or "bill of lading" in name.lower() or "draft bl" in name.lower()), None)
    if not si_name or not bl_name:
        result.update(status="NEEDS_REVIEW", review_reason="missing_attachment")
        return result
    if not attachment_texts[si_name].strip() or not attachment_texts[bl_name].strip():
        result.update(status="NEEDS_REVIEW", review_reason="unreadable_document")
        return result
    si_data, bl_data = extract_fields_from_doc(attachment_texts[si_name]), extract_fields_from_doc(attachment_texts[bl_name])
    for document in (si_data, bl_data):
        notify = str(document.get("notify_party") or "").upper()
        if not notify or "SAME AS" in notify:
            document["notify_party"] = document.get("consignee")
    result.update(si_data=si_data, bl_data=bl_data)
    if any(si_data[field] is None or bl_data[field] is None for field in FIELDS):
        result.update(status="NEEDS_REVIEW", review_reason="unreadable_document")
        return result
    defects = compare_fields(si_data, bl_data)
    result.update(status="MISMATCH" if defects else "OK", defect_fields=defects)
    return result


def inspect_email(email_record: dict, inbox: Inbox) -> dict:
    texts = {path: extract_text_from_bytes(inbox.read_bytes(path), path) for path in email_record.get("attachments", [])}
    return process_email(email_record, texts)


if __name__ == "__main__":
    inbox = Inbox(".")
    output = {email["email_id"]: inspect_email(email, inbox) for email in inbox}
    Path("submission.json").write_text(json.dumps(output, indent=2), encoding="utf-8")

